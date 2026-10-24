import csv
import io
from zipfile import BadZipFile, ZipFile
from django.contrib.auth.password_validation import validate_password
from django.core.exceptions import ValidationError
from django.db import transaction
from .models import User, UserRole, Classroom
from .services import set_membership, audit
from .permissions import require_admin

HEADERS = ['username', 'full_name', 'student_no', 'class_code']


def preview_import(upload):
    if upload.size > 2 * 1024 * 1024:
        raise ValidationError('导入文件不能超过 2 MB。')
    if upload.name.lower().endswith('.xlsx'):
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException

        try:
            with ZipFile(upload) as archive:
                if (
                    len(archive.infolist()) > 1000
                    or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024
                ):
                    raise ValidationError('Excel 文件展开后过大，请拆分名单再导入。')
            upload.seek(0)
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except (BadZipFile, InvalidFileException, ValueError, KeyError) as exc:
            raise ValidationError('无法读取 Excel 文件，请使用有效的 .xlsx 工作簿。') from exc
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header = [str(v or '').strip() for v in next(iterator, [])]
        if header != HEADERS:
            workbook.close()
            raise ValidationError('表头必须为：' + ', '.join(HEADERS))
        raw = [dict(zip(header, row)) for _, row in zip(range(501), iterator)]
        workbook.close()
    elif upload.name.lower().endswith('.csv'):
        try:
            reader = csv.DictReader(io.StringIO(upload.read().decode('utf-8-sig')))
            if reader.fieldnames != HEADERS:
                raise ValidationError('CSV 表头必须为：' + ', '.join(HEADERS))
            raw = [row for _, row in zip(range(501), reader)]
        except UnicodeDecodeError:
            raise ValidationError('CSV 必须采用 UTF-8 编码。')
    else:
        raise ValidationError('请使用 .csv 或 .xlsx 文件。')
    if not raw or len(raw) > 500:
        raise ValidationError('每批须为 1～500 行。')
    rows = []
    usernames, numbers = set(), set()
    for index, row in enumerate(raw, 2):
        row = {key: str(row.get(key) or '').strip() for key in HEADERS}
        error = ''
        if not all(row.values()):
            error = '所有列均需填写'
        elif row['username'] in usernames or row['student_no'] in numbers:
            error = '文件中存在重复账号或学号'
        elif (
            len(row['username']) > 150 or len(row['full_name']) > 80 or len(row['student_no']) > 40
        ):
            error = '字段长度超出限制'
        elif not Classroom.objects.filter(code=row['class_code'], archived=False).exists():
            error = '教学班不存在或已归档'
        else:
            user = User.objects.filter(username=row['username']).first()
            if user and (
                user.student_no != row['student_no']
                or user.full_name != row['full_name']
                or not user.has_role('student')
            ):
                error = '账号已存在且信息不一致，不自动覆盖'
            elif (
                User.objects.filter(student_no=row['student_no'])
                .exclude(username=row['username'])
                .exists()
            ):
                error = '学号已被其他账号使用'
        usernames.add(row['username'])
        numbers.add(row['student_no'])
        rows.append({**row, 'line': index, 'error': error})
    return rows


@transaction.atomic
def commit_import(actor, rows, password, *, allow_default_password=False):
    require_admin(actor)
    User.objects.select_for_update().get(pk=actor.pk)
    if not (allow_default_password and password == '123456'):
        validate_password(password)
    created = 0
    for row in rows:
        if row.get('error'):
            raise ValidationError('请修正所有错误后重新导入。')
        classroom = Classroom.objects.select_for_update().get(
            code=row['class_code'], archived=False
        )
        user = User.objects.filter(username=row['username']).first()
        if user:
            if (
                user.student_no != row['student_no']
                or user.full_name != row['full_name']
                or not user.has_role('student')
            ):
                raise ValidationError('预览后账号状态已变化，请重新预览。')
        else:
            if User.objects.filter(student_no=row['student_no']).exists():
                raise ValidationError('学号已存在，请重新预览。')
            user = User.objects.create_user(
                username=row['username'],
                full_name=row['full_name'],
                student_no=row['student_no'],
                password=password,
            )
            UserRole.objects.create(user=user, role='student', granted_by=actor)
            created += 1
        set_membership(actor, classroom, user)
    audit(actor, 'users.import', 'users', {'rows': len(rows), 'created': created})
    return created
